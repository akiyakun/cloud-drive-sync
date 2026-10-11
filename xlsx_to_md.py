"""XLSX をシートごとの Markdown に変換する（AI が直接読めない XLSX を読める形にするため）

- セル番地（A1 など）付きの表にして、セルで作った図の位置関係も追えるようにする
- セルの背景色を〔色A〕、文字の色を〔文字色A〕などの記号で残す（凡例の色分けや、目立たせたい一言など、色に意味がある資料のため）
  セルの中の一部の文字だけに色が付いている場合は〔文字色A:その文字〕の形にする
- 取り消し線を残す（「やめた案」を表すことが多いため）
- 結合したセルは〔結合 A4:C4〕の印を付ける
- 画像・図形描画・グラフは xlsx_drawings.py で読み取る
- 条件付き書式・入力規則（プルダウンなど）・数式・メモ／コメントを書き出す
"""

import io
import posixpath
import re
import warnings
import zipfile
from collections import Counter
from datetime import date, datetime, time

from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string

from colors import ColorLabels, color_name, is_notable_text_color, normalize
from xlsx_drawings import describe_anchors, image_files, read_sheet_drawings

# スレッド形式コメントの先頭に付く「[Threaded comment] ... Comment:」までの定型文
THREADED_HEADER = re.compile(r"^\[Threaded comment\].*?\nComment:\n", re.S)
HYPERLINK_FORMULA = re.compile(r'^=HYPERLINK\(\s*"([^"]*)"', re.I)
IMAGE_FORMULA = re.compile(r'^=IMAGE\(\s*"([^"]*)"', re.I)

CF_OPERATORS = {
    "greaterThan": "より大きい", "greaterThanOrEqual": "以上", "lessThan": "より小さい",
    "lessThanOrEqual": "以下", "equal": "と等しい", "notEqual": "と等しくない",
}


def _comment_text(comment):
    """セルのコメントから本文だけを取り出す

    Google の「コメント」は Excel のスレッド形式で書き出され、先頭に Excel の英語の定型文が付くため取り除く
    """
    text = comment.text.replace("\r\n", "\n")
    if text.startswith("[Threaded comment]"):
        text = THREADED_HEADER.sub("", text)
        text = text.replace("\nReply:\n", "\n返信: ")
    lines = [line.strip() for line in text.strip().split("\n")]
    return "<br>".join(line for line in lines if line)


def _rgb_of(color):
    """openpyxl の色（テーマ色などは扱わない）を #RRGGBB にする"""
    if color is None or getattr(color, "type", None) != "rgb":
        return None
    return normalize(color.rgb)


def _escape(text):
    return text.replace("\r\n", "\n").replace("|", "\\|").replace("\n", "<br>")


def _plain_value(value):
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # 10.0 → 10
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


class SheetContext:
    """シート 1 枚分の変換中に使う情報（色の記号と、色ごとの使用数）"""

    def __init__(self, labels):
        self.labels = labels
        self.fill_counts = Counter()
        self.font_counts = Counter()

    def font_mark(self, color, text=None):
        self.font_counts[color] += 1
        label = self.labels.label(color)
        return f"〔文字色{label}:{text}〕" if text is not None else f"〔文字色{label}〕"


def _rich_text(value, ctx):
    """セルの中の一部の文字だけに付いた色・取り消し線を残す"""
    parts = []
    for block in value:
        if isinstance(block, str):
            parts.append(_escape(block))
            continue
        text = _escape(block.text)
        font = block.font
        if font is not None and font.strike:
            text = f"~~{text}~~"
        color = _rgb_of(font.color) if font is not None else None
        if color and is_notable_text_color(color):
            text = ctx.font_mark(color, text)
        parts.append(text)
    return "".join(parts)


def _cell_text(cell, formula, ctx):
    """セル 1 つ分の表示。数式の HYPERLINK / IMAGE はリンク・画像として出す"""
    image = IMAGE_FORMULA.match(formula) if formula else None
    if image:
        # IMAGE 関数の画像は XLSX に入らない（値が #REF! になる）ので、URL だけ残す
        return f"![IMAGE 関数の画像](<{image.group(1)}>)"

    value = cell.value
    if value is None:
        return ""
    if isinstance(value, CellRichText):
        text = _rich_text(value, ctx)
        whole_cell_color = None
    else:
        text = _escape(_plain_value(value))
        if cell.font is not None and cell.font.strike:
            text = f"~~{text}~~"
        whole_cell_color = _rgb_of(cell.font.color) if cell.font is not None else None

    link = HYPERLINK_FORMULA.match(formula) if formula else None
    target = link.group(1) if link else (cell.hyperlink.target if cell.hyperlink is not None else None)
    if target:
        return f"[{text}](<{target}>)"  # リンクの青い文字は色の印を付けない
    if whole_cell_color and is_notable_text_color(whole_cell_color):
        text = f"{text} {ctx.font_mark(whole_cell_color)}"
    return text


def _fill_color(cell):
    """セルの背景色。塗りなし・白は None"""
    fill = cell.fill
    if fill is None or fill.fill_type != "solid":
        return None
    color = _rgb_of(fill.fgColor)
    return None if color in (None, "#FFFFFF") else color


def position_key(coord):
    """セル番地を「上から下、左から右」の順に並べるためのキー。位置不明は最後"""
    if coord is None:
        return (float("inf"), float("inf"))
    col, row = coordinate_from_string(coord)
    return (row, column_index_from_string(col))


def _conditional_formats(ws):
    lines = []
    for cf in ws.conditional_formatting:
        for rule in cf.rules:
            if rule.type == "cellIs":
                values = " と ".join(rule.formula or [])
                if rule.operator in ("between", "notBetween"):
                    cond = f"セルの値が {values} の" + ("間" if rule.operator == "between" else "範囲外")
                else:
                    cond = f"セルの値が {values} {CF_OPERATORS.get(rule.operator, rule.operator)}"
            elif rule.type == "expression":
                cond = f"数式 ={(rule.formula or [''])[0]} が成り立つ"
            elif rule.type in ("containsText", "notContainsText", "beginsWith", "endsWith"):
                verb = {"containsText": "を含む", "notContainsText": "を含まない",
                        "beginsWith": "で始まる", "endsWith": "で終わる"}[rule.type]
                cond = f"文字列「{rule.text}」{verb}"
            elif rule.type == "containsBlanks":
                cond = "空白"
            elif rule.type == "notContainsBlanks":
                cond = "空白でない"
            elif rule.type == "colorScale":
                cond = "カラースケール（値の大きさに応じて色が変わる）"
            elif rule.type == "dataBar":
                cond = "データバー"
            else:
                cond = rule.type

            effects = []
            dxf = rule.dxf
            if dxf is not None and dxf.fill is not None:
                bg = _rgb_of(dxf.fill.fgColor) or _rgb_of(dxf.fill.bgColor)
                if bg:
                    effects.append(f"背景 {bg}（{color_name(bg)}）")
            if dxf is not None and dxf.font is not None:
                fg = _rgb_of(dxf.font.color)
                if fg:
                    effects.append(f"文字 {fg}（{color_name(fg)}）")
                if dxf.font.strike:
                    effects.append("取り消し線")
                if dxf.font.b:
                    effects.append("太字")
            effect = "、".join(effects) if effects else "書式の変更"
            lines.append(f"- {cf.sqref}: {cond} → {effect}")
    return lines


def _data_validations(ws):
    lines = []
    for dv in ws.data_validations.dataValidation:
        formula = dv.formula1 or ""
        if dv.type == "list" and formula.startswith('"'):
            choices = " / ".join(c.strip() for c in formula.strip('"').split(","))
            text = f"プルダウン（{choices}）"
        elif dv.type == "list":
            text = f"プルダウン（{formula.replace('$', '')} の値から選ぶ）"
        else:
            cond = " ".join(x for x in (dv.operator, dv.formula1, dv.formula2) if x)
            text = f"入力規則 {dv.type}" + (f"（{cond}）" if cond else "")
        lines.append(f"- {dv.sqref}: {text}")
    return lines


def _sheet_markdown(ws, ws_formulas, anchors, images_dirname, labels):
    """1 シート分の Markdown を作る。labels は色の記号で、ファイル全体で共有する"""
    lines = [f"## {ws.title}"]
    if ws.sheet_state != "visible":
        lines.append("\n（非表示のシート）")

    formulas = {}
    for row in ws_formulas.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                formulas[c.coordinate] = c.value

    used = [(c.row, c.column) for row in ws.iter_rows() for c in row if c.value is not None]
    used += [position_key(coord) for coord in formulas]  # IMAGE 関数のセルは値が空のこともある
    if used:
        min_row, max_row = min(r for r, _ in used), max(r for r, _ in used)
        min_col, max_col = min(c for _, c in used), max(c for _, c in used)
        rows, cols = range(min_row, max_row + 1), range(min_col, max_col + 1)
        lines.append(f"\n範囲: {get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}")

        ctx = SheetContext(labels)
        positions = [(r, c) for r in rows for c in cols]
        fills = {pos: _fill_color(ws.cell(row=pos[0], column=pos[1])) for pos in positions}

        # 範囲の半分以上を占める背景色はシートの「背景」とみなし、セルごとには付けない
        fill_counts = Counter(color for color in fills.values() if color)
        background = None
        if fill_counts:
            top_color, top_count = fill_counts.most_common(1)[0]
            if top_count >= len(fills) / 2:
                background = top_color
                del fill_counts[background]

        merges = {str(m.coord).split(":")[0]: str(m.coord) for m in ws.merged_cells.ranges}

        table = []
        for r in rows:
            cells = []
            for c in cols:
                coord = f"{get_column_letter(c)}{r}"
                parts = [_cell_text(ws.cell(row=r, column=c), formulas.get(coord), ctx)]
                if fills[(r, c)] and fills[(r, c)] != background:
                    parts.append(f"〔色{labels.label(fills[(r, c)])}〕")
                if coord in merges:
                    parts.append(f"〔結合 {merges[coord]}〕")
                cells.append(" ".join(p for p in parts if p))
            table.append(f"| {r} | " + " | ".join(cells) + " |")

        if background:
            lines.append(f"\nシートの背景色: {background}（{color_name(background)}）")
        used_colors = sorted(set(fill_counts) | set(ctx.font_counts), key=labels.sort_key)
        if used_colors:
            lines.append("\n色の対応表（〔色A〕= セルの背景色、〔文字色A〕= 文字の色）:\n")
            for color in used_colors:
                usage = []
                if fill_counts[color]:
                    usage.append(f"背景 {fill_counts[color]} セル")
                if ctx.font_counts[color]:
                    usage.append(f"文字 {ctx.font_counts[color]} か所")
                lines.append(labels.legend_line(color, "、".join(usage)))

        lines.append("")
        lines.append("| | " + " | ".join(get_column_letter(c) for c in cols) + " |")
        lines.append("|---|" + "---|" * len(cols))
        lines.extend(table)
    else:
        lines.append("\n（セルに値なし）")

    lines.extend(describe_anchors(anchors, images_dirname, position_key))

    sections = [
        ("条件付き書式", _conditional_formats(ws)),
        ("入力規則", _data_validations(ws)),
        ("数式（表には計算結果を表示）", [
            f"- {coord}: `{formula}`" for coord, formula in sorted(formulas.items(), key=lambda kv: position_key(kv[0]))
        ]),
        ("セルのメモ・コメント", [
            f"- {c.coordinate}: {_comment_text(c.comment)}"
            for row in ws.iter_rows() for c in row if c.comment is not None
        ]),
    ]
    for title, items in sections:
        if items:
            lines.append(f"\n### {title}\n")
            lines.extend(items)

    return "\n".join(lines)


def xlsx_to_markdown(data, title, images_dirname):
    """XLSX のバイト列を Markdown に変換する

    戻り値: (Markdown 文字列, [(画像ファイル名, 画像データ), ...])
    """
    zf = zipfile.ZipFile(io.BytesIO(data))
    drawings = read_sheet_drawings(zf)

    with warnings.catch_warnings():
        # Google が書き出す XLSX には openpyxl 未対応の拡張が含まれ、警告が大量に出るため抑える
        warnings.simplefilter("ignore")
        wb = load_workbook(io.BytesIO(data), data_only=True, rich_text=True)  # 計算結果と文字ごとの書式
        wb_formulas = load_workbook(io.BytesIO(data), data_only=False)       # 数式そのもの

    parts = [
        f"# {title}",
        "",
        "> XLSX から自動生成。セル番地（A1 など）付き。"
        "セルの背景色は〔色A〕、文字の色は〔文字色A〕、一部の文字だけの色は〔文字色A:その文字〕のような記号で示す"
        "（同じアルファベットはファイル内で同じ色）。取り消し線は ~~ ~~、結合したセルは〔結合 A1:B2〕で示す。",
    ]
    labels = ColorLabels()
    for ws in wb.worksheets:
        parts.append("")
        parts.append(_sheet_markdown(ws, wb_formulas[ws.title], drawings.get(ws.title, []), images_dirname, labels))

    images = [(posixpath.basename(m), zf.read(m)) for m in image_files(drawings)]
    return "\n".join(parts) + "\n", images
