"""XLSX のシート上にある描画物（画像・図形描画・グラフ）を読み取り、Markdown 用の文章にする

- 画像: 取り出して、どのセルの辺りにあるかを記録する
- 図形描画: 画像ではなく図形として入っているので、図形の種類・文字・色と、矢印のつながり（「開始」→「終了」）を文章にする
- グラフ: 種類・タイトル・使っているデータの範囲を書き出す
"""

import math
import posixpath
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field

from openpyxl.utils import get_column_letter

from colors import color_name, normalize

NS = {
    "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
}
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
R_ID = f"{{{R_NS}}}id"
R_EMBED = f"{{{R_NS}}}embed"

# 線として扱う図形（矢印のつながりを調べる対象）
LINE_SHAPES = {"line", "straightConnector1", "bentConnector2", "bentConnector3", "bentConnector4",
               "bentConnector5", "curvedConnector2", "curvedConnector3", "curvedConnector4", "curvedConnector5"}

SHAPE_NAMES = {
    "rect": "四角形", "roundRect": "角丸四角形", "ellipse": "楕円", "diamond": "ひし形",
    "triangle": "三角形", "rtTriangle": "直角三角形", "parallelogram": "平行四辺形", "trapezoid": "台形",
    "pentagon": "五角形", "hexagon": "六角形", "octagon": "八角形", "star5": "星形", "heart": "ハート",
    "cloud": "雲", "can": "円柱", "cube": "立方体", "plus": "十字",
    "rightArrow": "右向きの矢印図形", "leftArrow": "左向きの矢印図形", "upArrow": "上向きの矢印図形",
    "downArrow": "下向きの矢印図形", "leftRightArrow": "左右の矢印図形", "upDownArrow": "上下の矢印図形",
    "wedgeRectCallout": "吹き出し", "wedgeRoundRectCallout": "吹き出し", "wedgeEllipseCallout": "吹き出し",
    "flowChartProcess": "処理（フローチャート）", "flowChartDecision": "判断（フローチャート）",
    "flowChartTerminator": "端子（フローチャート）", "flowChartInputOutput": "データ（フローチャート）",
    "flowChartDocument": "書類（フローチャート）", "flowChartPredefinedProcess": "定義済み処理（フローチャート）",
}

CHART_TYPES = {
    "barChart": "棒グラフ", "bar3DChart": "3D 棒グラフ", "lineChart": "折れ線グラフ", "line3DChart": "3D 折れ線グラフ",
    "pieChart": "円グラフ", "pie3DChart": "3D 円グラフ", "doughnutChart": "ドーナツグラフ", "areaChart": "面グラフ",
    "area3DChart": "3D 面グラフ", "scatterChart": "散布図", "radarChart": "レーダーチャート", "bubbleChart": "バブルチャート",
    "stockChart": "株価チャート", "surfaceChart": "等高線グラフ",
}


@dataclass
class Shape:
    id: str
    kind: str
    text: str
    fill: str | None
    box: tuple  # (x, y, 幅, 高さ)。同じ描画の中で比べられる座標（EMU）


@dataclass
class Line:
    kind: str
    start: tuple
    end: tuple
    arrow_at_start: bool
    arrow_at_end: bool
    start_id: str | None
    end_id: str | None


@dataclass
class Anchor:
    """シート上の 1 つの描画物（画像 1 枚、図形描画 1 つ、グラフ 1 つなど）"""
    start: str | None
    end: str | None
    images: list = field(default_factory=list)
    shapes: list = field(default_factory=list)
    lines: list = field(default_factory=list)
    charts: list = field(default_factory=list)


def rels(zf, part):
    """XLSX 内のファイル（例: xl/workbook.xml）が参照する先を {Id: ZIP 内のパス} で返す"""
    base, name = posixpath.split(part)
    rels_path = posixpath.join(base, "_rels", f"{name}.rels")
    if rels_path not in zf.namelist():
        return {}
    result = {}
    for r in ET.fromstring(zf.read(rels_path)).findall("rel:Relationship", NS):
        target = r.get("Target")
        if r.get("TargetMode") == "External":
            continue
        path = target.lstrip("/") if target.startswith("/") else posixpath.join(base, target)
        result[r.get("Id")] = posixpath.normpath(path)
    return result


def _marker_cell(marker):
    """描画物の位置情報（0 始まりの列・行）をセル番地にする"""
    if marker is None:
        return None
    col = int(marker.find("xdr:col", NS).text)
    row = int(marker.find("xdr:row", NS).text)
    return f"{get_column_letter(col + 1)}{row + 1}"


# --- 図形描画 ---------------------------------------------------------------

def _xfrm(sp_pr):
    """図形の位置と大きさ、反転の有無"""
    xfrm = sp_pr.find("a:xfrm", NS) if sp_pr is not None else None
    if xfrm is None:
        return (0, 0, 0, 0), False, False
    off, ext = xfrm.find("a:off", NS), xfrm.find("a:ext", NS)
    box = (
        int(off.get("x", 0)) if off is not None else 0,
        int(off.get("y", 0)) if off is not None else 0,
        int(ext.get("cx", 0)) if ext is not None else 0,
        int(ext.get("cy", 0)) if ext is not None else 0,
    )
    return box, xfrm.get("flipH") == "1", xfrm.get("flipV") == "1"


def _transform_box(box, t):
    """グループの中の座標を、描画全体の座標に直す。t = (拡大率x, 拡大率y, ずらしx, ずらしy)"""
    sx, sy, tx, ty = t
    x, y, w, h = box
    return (x * sx + tx, y * sy + ty, w * sx, h * sy)


def _group_transform(grp, parent):
    """グループの「子の座標系」を親の座標系に直すための変換を作る"""
    xfrm = grp.find("xdr:grpSpPr/a:xfrm", NS)
    if xfrm is None:
        return parent
    off, ext = xfrm.find("a:off", NS), xfrm.find("a:ext", NS)
    ch_off, ch_ext = xfrm.find("a:chOff", NS), xfrm.find("a:chExt", NS)
    if None in (off, ext, ch_off, ch_ext):
        return parent
    sx = int(ext.get("cx")) / int(ch_ext.get("cx")) if int(ch_ext.get("cx")) else 1
    sy = int(ext.get("cy")) / int(ch_ext.get("cy")) if int(ch_ext.get("cy")) else 1
    tx = int(off.get("x")) - int(ch_off.get("x")) * sx
    ty = int(off.get("y")) - int(ch_off.get("y")) * sy
    psx, psy, ptx, pty = parent
    return (sx * psx, sy * psy, tx * psx + ptx, ty * psy + pty)


def _shape_text(element):
    paragraphs = []
    for p in element.findall(".//a:p", NS):
        text = "".join(t.text or "" for t in p.findall(".//a:t", NS)).strip()
        if text:
            paragraphs.append(text)
    return " / ".join(paragraphs)


def _arrow(sp_pr, tag):
    end = sp_pr.find(f"a:ln/a:{tag}", NS) if sp_pr is not None else None
    return end is not None and end.get("type", "none") != "none"


def _read_line(element, kind, t):
    sp_pr = element.find("xdr:spPr", NS)
    box, flip_h, flip_v = _xfrm(sp_pr)
    x, y, w, h = _transform_box(box, t)
    x1, x2 = (x + w, x) if flip_h else (x, x + w)
    y1, y2 = (y + h, y) if flip_v else (y, y + h)
    st = element.find("xdr:nvCxnSpPr/xdr:cNvCxnSpPr/a:stCxn", NS)
    en = element.find("xdr:nvCxnSpPr/xdr:cNvCxnSpPr/a:endCxn", NS)
    return Line(
        kind=kind, start=(x1, y1), end=(x2, y2),
        arrow_at_start=_arrow(sp_pr, "headEnd"), arrow_at_end=_arrow(sp_pr, "tailEnd"),
        start_id=st.get("id") if st is not None else None,
        end_id=en.get("id") if en is not None else None,
    )


def _collect(element, anchor, t, drawing_rels, zf):
    """描画物の中身（グループの中も含む）を集める"""
    for child in element:
        tag = child.tag.split("}")[1]
        if tag == "grpSp":
            _collect(child, anchor, _group_transform(child, t), drawing_rels, zf)
        elif tag in ("sp", "cxnSp"):
            sp_pr = child.find("xdr:spPr", NS)
            geom = sp_pr.find("a:prstGeom", NS) if sp_pr is not None else None
            kind = geom.get("prst") if geom is not None else "custom"
            if tag == "cxnSp" or kind in LINE_SHAPES:
                anchor.lines.append(_read_line(child, kind, t))
                continue
            nv = child.find("xdr:nvSpPr/xdr:cNvPr", NS)
            is_textbox = child.find("xdr:nvSpPr/xdr:cNvSpPr[@txBox='1']", NS) is not None
            fill = sp_pr.find("a:solidFill/a:srgbClr", NS) if sp_pr is not None else None
            box, _, _ = _xfrm(sp_pr)
            anchor.shapes.append(Shape(
                id=nv.get("id") if nv is not None else "",
                kind="テキストボックス" if is_textbox else SHAPE_NAMES.get(kind, f"図形（{kind}）"),
                text=_shape_text(child),
                fill=normalize(fill.get("val")) if fill is not None else None,
                box=_transform_box(box, t),
            ))
        elif tag == "pic":
            blip = child.find(".//a:blip", NS)
            media = drawing_rels.get(blip.get(R_EMBED)) if blip is not None else None
            if media in zf.namelist():
                anchor.images.append(media)
        elif tag == "graphicFrame":
            ref = child.find(".//c:chart", NS)
            chart_part = drawing_rels.get(ref.get(R_ID)) if ref is not None else None
            if chart_part in zf.namelist():
                anchor.charts.append(_read_chart(zf, chart_part))


def read_sheet_drawings(zf):
    """シート名 → [Anchor, ...]"""
    names = set(zf.namelist())
    wb_rels = rels(zf, "xl/workbook.xml")
    result = {}
    for sheet in ET.fromstring(zf.read("xl/workbook.xml")).findall("main:sheets/main:sheet", NS):
        sheet_part = wb_rels.get(sheet.get(R_ID))
        if sheet_part not in names:
            continue
        sheet_rels = rels(zf, sheet_part)
        anchors = []
        for drawing in ET.fromstring(zf.read(sheet_part)).findall("main:drawing", NS):
            drawing_part = sheet_rels.get(drawing.get(R_ID))
            if drawing_part not in names:
                continue
            drawing_rels = rels(zf, drawing_part)
            for element in ET.fromstring(zf.read(drawing_part)):
                anchor = Anchor(
                    start=_marker_cell(element.find("xdr:from", NS)),
                    end=_marker_cell(element.find("xdr:to", NS)),
                )
                _collect(element, anchor, (1, 1, 0, 0), drawing_rels, zf)
                if anchor.images or anchor.shapes or anchor.lines or anchor.charts:
                    anchors.append(anchor)
        result[sheet.get("name")] = anchors
    return result


def _distance_to_box(point, box):
    px, py = point
    x, y, w, h = box
    dx = max(x - px, 0, px - (x + w))
    dy = max(y - py, 0, py - (y + h))
    return math.hypot(dx, dy)


def _nearest_shape(point, shapes):
    """線の端に一番近い図形。離れすぎていたら None（つながっていないとみなす）"""
    best = None
    for shape in shapes:
        d = _distance_to_box(point, shape.box)
        tolerance = 0.15 * max(shape.box[2], shape.box[3]) + 12700  # 図形の大きさの 15% + 1pt
        if d <= tolerance and (best is None or d < best[0]):
            best = (d, shape)
    return best[1] if best else None


def describe_shapes(anchor):
    """図形描画 1 つ分を Markdown の箇条書きにする"""
    shapes = sorted(anchor.shapes, key=lambda s: (s.box[1], s.box[0]))  # 上から下、左から右
    number = {id(s): i + 1 for i, s in enumerate(shapes)}
    by_id = {s.id: s for s in shapes}

    def label(shape):
        if shape is None:
            return "（つながり先なし）"
        return f"図形{number[id(shape)]}" + (f"「{shape.text}」" if shape.text else "")

    lines = []
    for s in shapes:
        detail = f"塗り: {s.fill} {color_name(s.fill)}" if s.fill else ""
        lines.append(f"  - 図形{number[id(s)]}: {s.kind}" + (f"「{s.text}」" if s.text else "")
                     + (f"（{detail}）" if detail else ""))
    for line in anchor.lines:
        a = by_id.get(line.start_id) or _nearest_shape(line.start, shapes)
        b = by_id.get(line.end_id) or _nearest_shape(line.end, shapes)
        if line.arrow_at_end and not line.arrow_at_start:
            text = f"{label(a)} → {label(b)}"
        elif line.arrow_at_start and not line.arrow_at_end:
            text = f"{label(b)} → {label(a)}"
        elif line.arrow_at_start and line.arrow_at_end:
            text = f"{label(a)} ↔ {label(b)}"
        else:
            text = f"{label(a)} — {label(b)}（矢印なしの線）"
        lines.append(f"  - 矢印: {text}" if (line.arrow_at_start or line.arrow_at_end) else f"  - 線: {text}")
    return lines


# --- グラフ -----------------------------------------------------------------

def _ref_or_value(element):
    """系列名などの「セル参照」か「値」を文字にする"""
    if element is None:
        return None
    f = element.find(".//c:f", NS)
    if f is not None and f.text:
        return f.text.replace("$", "")
    v = element.find(".//c:v", NS)
    return v.text if v is not None else None


def _read_chart(zf, chart_part):
    root = ET.fromstring(zf.read(chart_part))
    chart = root.find("c:chart", NS)
    title_el = chart.find("c:title", NS) if chart is not None else None
    title = "".join(t.text or "" for t in title_el.findall(".//a:t", NS)) if title_el is not None else ""
    plot = chart.find("c:plotArea", NS) if chart is not None else None
    kinds, series = [], []
    for child in (plot if plot is not None else []):
        tag = child.tag.split("}")[1]
        if tag not in CHART_TYPES:
            continue
        kind = CHART_TYPES[tag]
        bar_dir = child.find("c:barDir", NS)
        if bar_dir is not None:
            kind = ("横" if bar_dir.get("val") == "bar" else "縦") + kind
        kinds.append(kind)
        for ser in child.findall("c:ser", NS):
            series.append({
                "name": _ref_or_value(ser.find("c:tx", NS)),
                "categories": _ref_or_value(ser.find("c:cat", NS)) or _ref_or_value(ser.find("c:xVal", NS)),
                "values": _ref_or_value(ser.find("c:val", NS)) or _ref_or_value(ser.find("c:yVal", NS)),
            })
    return {"title": title, "kinds": kinds, "series": series}


def describe_chart(chart):
    lines = []
    for i, s in enumerate(chart["series"], 1):
        parts = []
        if s["name"]:
            parts.append(f"名前 {s['name']}")
        if s["values"]:
            parts.append(f"値 {s['values']}")
        if s["categories"]:
            parts.append(f"項目 {s['categories']}")
        lines.append(f"  - 系列{i}: " + "、".join(parts))
    return lines


def describe_anchors(anchors, images_dirname, position_key):
    """シート 1 枚分の描画物を Markdown の節にする"""
    out = []
    ordered = sorted(anchors, key=lambda a: position_key(a.start))

    def where(a):
        if a.start and a.end:
            return f"{a.start}〜{a.end}"
        return f"{a.start} 付近" if a.start else "位置不明"

    images = [(a, m) for a in ordered for m in a.images]
    if images:
        out.append("\n### 画像（シート上の位置順）\n")
        for a, media in images:
            filename = posixpath.basename(media)
            out.append(f"- {where(a)}: ![{filename}](<{images_dirname}/{filename}>)")

    drawings = [a for a in ordered if a.shapes or a.lines]
    if drawings:
        out.append("\n### 図形描画\n")
        for a in drawings:
            out.append(f"- {where(a)}:")
            out.extend(describe_shapes(a))

    charts = [(a, c) for a in ordered for c in a.charts]
    if charts:
        out.append("\n### グラフ\n")
        for a, c in charts:
            kind = "・".join(c["kinds"]) or "グラフ"
            out.append(f"- {where(a)}: {kind}" + (f"「{c['title']}」" if c["title"] else ""))
            out.extend(describe_chart(c))
    return out


def image_files(drawings_by_sheet):
    """取り出す画像の ZIP 内パス（同じ画像が何か所にあっても 1 つ）"""
    return sorted({m for anchors in drawings_by_sheet.values() for a in anchors for m in a.images})
